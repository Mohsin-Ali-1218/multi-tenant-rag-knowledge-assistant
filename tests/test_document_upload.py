import io
import uuid
import zipfile
import logging
from unittest.mock import patch, AsyncMock, MagicMock
import pytest
import pypdf
import docx
from bson import ObjectId
from httpx import AsyncClient, ASGITransport
from backend.main import app
from backend.app.config import settings
from backend.app.db.mongo import get_db
from backend.app.db.vector import search_tenant_docs
from backend.app.services.embedding import get_embedding
import backend.app.api.documents

def build_pdf_bytes(pages_text: list[str]) -> bytes:
    """Builds a valid multi-page PDF with extractable text."""
    lines = [b'%PDF-1.4\n']
    offsets = {}

    def add_obj(obj_num, data):
        offsets[obj_num] = sum(len(l) for l in lines)
        lines.append(f'{obj_num} 0 obj\n'.encode() + data + b'\nendobj\n')

    total_pages = len(pages_text)
    page_obj_ids = [3 + i * 3 + 2 for i in range(total_pages)]

    add_obj(1, b'<< /Type /Catalog /Pages 2 0 R >>')
    kids_str = ' '.join(f'{pid} 0 R' for pid in page_obj_ids)
    add_obj(2, f'<< /Type /Pages /Kids [{kids_str}] /Count {total_pages} >>'.encode())

    curr = 3
    for text in pages_text:
        cid = curr
        fid = curr + 1
        pid = curr + 2
        curr += 3
        stream_data = f'BT /F1 12 Tf 50 700 Td ({text}) Tj ET'.encode()
        add_obj(cid, f'<< /Length {len(stream_data)} >> stream\n'.encode() + stream_data + b'\nendstream')
        add_obj(fid, b'<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>')
        add_obj(pid, f'<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents {cid} 0 R /Resources << /Font << /F1 {fid} 0 R >> >> >>'.encode())

    startxref = sum(len(l) for l in lines)
    total_objs = curr
    lines.append(f'xref\n0 {total_objs}\n0000000000 65535 f \n'.encode())
    for i in range(1, total_objs):
        lines.append(f'{offsets[i]:010d} 00000 n \n'.encode())
    lines.append(f'trailer << /Root 1 0 R /Size {total_objs} >>\nstartxref\n{startxref}\n%%EOF\n'.encode())
    return b''.join(lines)

def build_blank_pdf_bytes() -> bytes:
    writer = pypdf.PdfWriter()
    writer.add_blank_page(width=100, height=100)
    buf = io.BytesIO()
    writer.write(buf)
    return buf.getvalue()

def build_encrypted_pdf_bytes() -> bytes:
    writer = pypdf.PdfWriter()
    writer.add_blank_page(width=100, height=100)
    writer.encrypt("secret123")
    buf = io.BytesIO()
    writer.write(buf)
    return buf.getvalue()

def build_docx_bytes(items: list) -> bytes:
    """Builds a valid DOCX with mixed paragraphs and tables in order."""
    doc = docx.Document()
    for item in items:
        if isinstance(item, str):
            doc.add_paragraph(item)
        elif isinstance(item, list):
            tbl = doc.add_table(rows=len(item), cols=len(item[0]))
            for r_idx, row in enumerate(item):
                for c_idx, cell in enumerate(row):
                    tbl.rows[r_idx].cells[c_idx].text = str(cell)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()

def build_non_ooxml_zip_bytes() -> bytes:
    """Valid ZIP file but missing word/document.xml OOXML structure."""
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as zf:
        zf.writestr("notes.txt", "This is just a plain zip file renamed to docx.")
    return buf.getvalue()

async def get_auth_token(ac: AsyncClient, email: str, tenant_name: str) -> tuple[str, str]:
    res = await ac.post("/auth/login", json={"email": email, "password": "password123"})
    if res.status_code == 401:
        res = await ac.post("/auth/register", json={
            "email": email,
            "password": "password123",
            "tenant_name": tenant_name
        })
    assert res.status_code == 200, f"Auth failed for {email}: {res.text}"
    data = res.json()
    return data["access_token"], data["tenant_id"]

@pytest.mark.asyncio
async def test_document_uploads_and_grounded_metadata():
    async with AsyncClient(transport=ASGITransport(app=app), base_url="http://testserver") as ac:
        pfx = uuid.uuid4().hex[:6]
        token_a, tenant_a = await get_auth_token(ac, f"m2_admin_a_{pfx}@tenant.com", f"M2 Tenant A {pfx}")
        headers_a = {"Authorization": f"Bearer {token_a}"}

        token_b, tenant_b = await get_auth_token(ac, f"m2_admin_b_{pfx}@tenant.com", f"M2 Tenant B {pfx}")
        headers_b = {"Authorization": f"Bearer {token_b}"}

        # ---------------------------------------------------------------------
        # 1. Multi-page PDF Upload & Page Numbering
        # ---------------------------------------------------------------------
        pdf_bytes = build_pdf_bytes([
            "Alpha Clause: All proprietary systems must remain encrypted at rest.",
            "Beta Clause: Confidential customer data must be retained for exactly seven years."
        ])
        res_pdf = await ac.post(
            "/documents/upload",
            headers=headers_a,
            files={"file": ("legal_policy.pdf", pdf_bytes, "application/pdf")},
            data={"title": "Legal Policy Manual"}
        )
        assert res_pdf.status_code == 201, f"PDF upload failed: {res_pdf.text}"
        data_pdf = res_pdf.json()
        assert data_pdf["title"] == "Legal Policy Manual"
        assert data_pdf["filename"] == "legal_policy.pdf"
        assert data_pdf["file_type"] == "pdf"
        assert data_pdf["chunk_count"] >= 2
        assert "Alpha Clause" in data_pdf["raw_text"]
        assert "Beta Clause" in data_pdf["raw_text"]

        # Verify Qdrant points received page numbers
        query_vec = get_embedding("Alpha Clause proprietary systems")
        hits_pdf = await search_tenant_docs(tenant_a, query_vec, limit=5, score_threshold=0.3)
        pdf_hits = [h for h in hits_pdf if h["document_id"] == data_pdf["id"]]
        assert len(pdf_hits) > 0
        assert pdf_hits[0]["page_number"] in [1, 2]
        assert pdf_hits[0]["filename"] == "legal_policy.pdf"

        # ---------------------------------------------------------------------
        # 2. DOCX Upload with Ordered Paragraphs & Tables
        # ---------------------------------------------------------------------
        docx_content = [
            "Section 1: Onboarding Protocol",
            [["Role", "Hardware"], ["Engineer", "MacBook Pro M3"], ["Analyst", "ThinkPad P16"]],
            "Section 2: Security Clearance Verification"
        ]
        docx_bytes = build_docx_bytes(docx_content)
        res_docx = await ac.post(
            "/documents/upload",
            headers=headers_a,
            files={"file": ("onboarding.docx", docx_bytes, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")}
        )
        assert res_docx.status_code == 201, f"DOCX upload failed: {res_docx.text}"
        data_docx = res_docx.json()
        assert data_docx["title"] == "onboarding"
        assert data_docx["filename"] == "onboarding.docx"
        assert data_docx["file_type"] == "docx"
        
        # Verify order: Section 1 -> Table -> Section 2
        sec1_idx = data_docx["raw_text"].find("Section 1")
        tbl_idx = data_docx["raw_text"].find("Engineer | MacBook Pro M3")
        sec2_idx = data_docx["raw_text"].find("Section 2")
        assert sec1_idx < tbl_idx < sec2_idx, "DOCX paragraph and table order was not preserved!"

        # Verify Qdrant points have page_number=None for DOCX
        query_vec_docx = get_embedding("Hardware MacBook Pro Engineer")
        hits_docx = await search_tenant_docs(tenant_a, query_vec_docx, limit=5, score_threshold=0.3)
        docx_hits = [h for h in hits_docx if h["document_id"] == data_docx["id"]]
        assert len(docx_hits) > 0
        assert docx_hits[0]["page_number"] is None
        assert docx_hits[0]["filename"] == "onboarding.docx"

        # ---------------------------------------------------------------------
        # 3. TXT Upload with UTF-8 & UTF-8 BOM
        # ---------------------------------------------------------------------
        txt_content = "\ufeffServer Configuration:\nHost=127.0.0.1\nPort=8080\nEnv=Production"
        res_txt = await ac.post(
            "/documents/upload",
            headers=headers_a,
            files={"file": ("config.txt", txt_content.encode("utf-8"), "text/plain")}
        )
        assert res_txt.status_code == 201, f"TXT upload failed: {res_txt.text}"
        data_txt = res_txt.json()
        assert data_txt["filename"] == "config.txt"
        assert data_txt["file_type"] == "txt"
        assert "\ufeff" not in data_txt["raw_text"], "UTF-8 BOM should be stripped"
        assert "Host=127.0.0.1" in data_txt["raw_text"]

        # ---------------------------------------------------------------------
        # 4. Raw-text Backwards Compatibility (POST /documents)
        # ---------------------------------------------------------------------
        res_raw = await ac.post(
            "/documents",
            headers=headers_a,
            json={
                "title": "Legacy Raw Document",
                "raw_text": "This document was submitted via the legacy JSON payload without files."
            }
        )
        assert res_raw.status_code == 201
        data_raw = res_raw.json()
        assert data_raw["title"] == "Legacy Raw Document"
        assert data_raw["filename"] is None
        assert data_raw["file_type"] == "raw"

        # ---------------------------------------------------------------------
        # 5. Validation & Edge Cases: Empty, Corrupt, Scanned, Encrypted
        # ---------------------------------------------------------------------
        # Empty file (0 bytes)
        res_empty = await ac.post(
            "/documents/upload",
            headers=headers_a,
            files={"file": ("empty.txt", b"", "text/plain")}
        )
        assert res_empty.status_code == 400
        assert "empty" in res_empty.json()["detail"].lower()

        # Whitespace-only text file
        res_ws = await ac.post(
            "/documents/upload",
            headers=headers_a,
            files={"file": ("blank.txt", b"   \n\t  \n  ", "text/plain")}
        )
        assert res_ws.status_code == 400

        # Scanned / Image-only PDF with 0 extractable characters -> HTTP 422
        blank_pdf = build_blank_pdf_bytes()
        res_scanned = await ac.post(
            "/documents/upload",
            headers=headers_a,
            files={"file": ("scanned_invoice.pdf", blank_pdf, "application/pdf")}
        )
        assert res_scanned.status_code == 422
        assert "scanned" in res_scanned.json()["detail"].lower()

        # Encrypted PDF -> HTTP 400
        encrypted_pdf = build_encrypted_pdf_bytes()
        res_enc = await ac.post(
            "/documents/upload",
            headers=headers_a,
            files={"file": ("encrypted.pdf", encrypted_pdf, "application/pdf")}
        )
        assert res_enc.status_code == 400
        assert "encrypted" in res_enc.json()["detail"].lower()

        # Corrupt PDF (garbage bytes with .pdf extension) -> HTTP 400
        res_corrupt_pdf = await ac.post(
            "/documents/upload",
            headers=headers_a,
            files={"file": ("corrupt.pdf", b"Not a PDF file at all", "application/pdf")}
        )
        assert res_corrupt_pdf.status_code == 400

        # Corrupt DOCX (garbage bytes) -> HTTP 400
        res_corrupt_docx = await ac.post(
            "/documents/upload",
            headers=headers_a,
            files={"file": ("corrupt.docx", b"Not a docx file", "application/vnd.openxmlformats-officedocument.wordprocessingml.document")}
        )
        assert res_corrupt_docx.status_code == 400

        # Fake DOCX (valid ZIP but missing word/document.xml OOXML structure) -> HTTP 400
        fake_docx = build_non_ooxml_zip_bytes()
        res_fake_docx = await ac.post(
            "/documents/upload",
            headers=headers_a,
            files={"file": ("fake.docx", fake_docx, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")}
        )
        assert res_fake_docx.status_code == 400
        assert "ooxml" in res_fake_docx.json()["detail"].lower()

        # Non-UTF-8 TXT (invalid encoding) -> HTTP 400
        invalid_bytes = b"Invalid UTF-8 \xff\xfe bytes"
        res_bad_encoding = await ac.post(
            "/documents/upload",
            headers=headers_a,
            files={"file": ("bad_encoding.txt", invalid_bytes, "text/plain")}
        )
        assert res_bad_encoding.status_code == 400
        assert "utf-8" in res_bad_encoding.json()["detail"].lower()

        # Unsupported format (e.g. .exe or .py) -> HTTP 400
        res_unsupported = await ac.post(
            "/documents/upload",
            headers=headers_a,
            files={"file": ("script.py", b"print('hello')", "text/x-python")}
        )
        assert res_unsupported.status_code == 400
        assert "unsupported" in res_unsupported.json()["detail"].lower()

        # ---------------------------------------------------------------------
        # 6. Resource Limits
        # ---------------------------------------------------------------------
        # File size exceeds limit (simulate limit check)
        oversized_bytes = b"X" * (settings.MAX_UPLOAD_SIZE_MB * 1024 * 1024 + 1024)
        res_oversized = await ac.post(
            "/documents/upload",
            headers=headers_a,
            files={"file": ("oversized.txt", oversized_bytes, "text/plain")}
        )
        assert res_oversized.status_code == 413

        # PDF page count exceeds limit
        many_pages = build_pdf_bytes(["Sample text"] * (settings.MAX_DOCUMENT_PAGES + 1))
        res_pages = await ac.post(
            "/documents/upload",
            headers=headers_a,
            files={"file": ("too_many_pages.pdf", many_pages, "application/pdf")}
        )
        assert res_pages.status_code == 400
        assert "page count" in res_pages.json()["detail"].lower()

        # ---------------------------------------------------------------------
        # 7. Unauthenticated Requests
        # ---------------------------------------------------------------------
        res_unauth = await ac.post(
            "/documents/upload",
            files={"file": ("test.txt", b"Hello", "text/plain")}
        )
        assert res_unauth.status_code == 401

        # ---------------------------------------------------------------------
        # 8. Cross-Tenant Isolation & Denial Behavior
        # ---------------------------------------------------------------------
        # Tenant B attempts to access Tenant A's uploaded PDF by ID -> 403 Forbidden
        doc_a_id = data_pdf["id"]
        res_cross_get = await ac.get(f"/documents/{doc_a_id}", headers=headers_b)
        assert res_cross_get.status_code == 403
        assert "forbidden" in res_cross_get.json()["detail"].lower()

        # Tenant B search should return 0 results from Tenant A's document
        hits_b = await search_tenant_docs(tenant_b, query_vec, limit=5, score_threshold=0.3)
        assert not any(h["document_id"] == doc_a_id for h in hits_b), "Cross-tenant retrieval leaked Tenant A document!"

@pytest.mark.asyncio
async def test_single_and_multipage_pdf_exact_page_metadata():
    async with AsyncClient(transport=ASGITransport(app=app), base_url="http://testserver") as ac:
        pfx = uuid.uuid4().hex[:6]
        token, tenant_id = await get_auth_token(ac, f"pdf_exact_{pfx}@tenant.com", f"PDF Exact Tenant {pfx}")
        headers = {"Authorization": f"Bearer {token}"}

        # 1. Single-page PDF with exact page_number == 1 and filename
        single_pdf = build_pdf_bytes(["Sole paragraph on single page document regarding distributed systems."])
        res_single = await ac.post(
            "/documents/upload",
            headers=headers,
            files={"file": ("single_doc.pdf", single_pdf, "application/pdf")}
        )
        assert res_single.status_code == 201
        data_single = res_single.json()
        assert data_single["filename"] == "single_doc.pdf"
        assert data_single["file_type"] == "pdf"

        q_vec = get_embedding("Sole paragraph on single page distributed systems")
        hits_single = await search_tenant_docs(tenant_id, q_vec, limit=3, score_threshold=0.3)
        single_hits = [h for h in hits_single if h["document_id"] == data_single["id"]]
        assert len(single_hits) > 0
        for hit in single_hits:
            assert hit["page_number"] == 1
            assert hit["filename"] == "single_doc.pdf"

        # 2. Multi-page PDF: explicit verification of both page 1 and page 2
        multipage_pdf = build_pdf_bytes([
            "Gamma Section: First page content regarding server architectures and clusters.",
            "Delta Section: Second page content regarding database migrations and sharding."
        ])
        res_multi = await ac.post(
            "/documents/upload",
            headers=headers,
            files={"file": ("multipage_doc.pdf", multipage_pdf, "application/pdf")}
        )
        assert res_multi.status_code == 201
        data_multi = res_multi.json()

        # Query for page 1 content
        q_vec_p1 = get_embedding("Gamma Section server architectures clusters")
        hits_p1 = await search_tenant_docs(tenant_id, q_vec_p1, limit=3, score_threshold=0.3)
        doc_hits_p1 = [h for h in hits_p1 if h["document_id"] == data_multi["id"]]
        assert len(doc_hits_p1) > 0
        assert doc_hits_p1[0]["page_number"] == 1
        assert doc_hits_p1[0]["filename"] == "multipage_doc.pdf"

        # Query for page 2 content
        q_vec_p2 = get_embedding("Delta Section database migrations sharding")
        hits_p2 = await search_tenant_docs(tenant_id, q_vec_p2, limit=3, score_threshold=0.3)
        doc_hits_p2 = [h for h in hits_p2 if h["document_id"] == data_multi["id"]]
        assert len(doc_hits_p2) > 0
        assert doc_hits_p2[0]["page_number"] == 2
        assert doc_hits_p2[0]["filename"] == "multipage_doc.pdf"

@pytest.mark.asyncio
async def test_pure_utf8_txt_and_raw_text_payload_metadata():
    async with AsyncClient(transport=ASGITransport(app=app), base_url="http://testserver") as ac:
        pfx = uuid.uuid4().hex[:6]
        token, tenant_id = await get_auth_token(ac, f"txt_exact_{pfx}@tenant.com", f"TXT Exact Tenant {pfx}")
        headers = {"Authorization": f"Bearer {token}"}

        # Pure UTF-8 without BOM
        pure_utf8_content = "Specification doc with UTF-8 characters: temperature ±5°C, price €100, trademark ™."
        res_txt = await ac.post(
            "/documents/upload",
            headers=headers,
            files={"file": ("pure_utf8.txt", pure_utf8_content.encode("utf-8"), "text/plain")}
        )
        assert res_txt.status_code == 201
        data_txt = res_txt.json()
        assert data_txt["filename"] == "pure_utf8.txt"
        assert data_txt["file_type"] == "txt"

        # Verify vector chunks have page_number is None and filename
        q_vec = get_embedding("Specification doc temperature price")
        hits_txt = await search_tenant_docs(tenant_id, q_vec, limit=3, score_threshold=0.3)
        matching_txt = [h for h in hits_txt if h["document_id"] == data_txt["id"]]
        assert len(matching_txt) > 0
        assert matching_txt[0]["page_number"] is None
        assert matching_txt[0]["filename"] == "pure_utf8.txt"

        # Raw-text legacy payload (POST /documents)
        res_raw = await ac.post(
            "/documents",
            headers=headers,
            json={
                "title": "Raw Text Document",
                "raw_text": "Content submitted directly as raw text JSON for vector indexing."
            }
        )
        assert res_raw.status_code == 201
        data_raw = res_raw.json()

        q_vec_raw = get_embedding("Content submitted directly as raw text")
        hits_raw = await search_tenant_docs(tenant_id, q_vec_raw, limit=3, score_threshold=0.3)
        matching_raw = [h for h in hits_raw if h["document_id"] == data_raw["id"]]
        assert len(matching_raw) > 0
        assert matching_raw[0]["page_number"] is None
        assert matching_raw[0]["filename"] is None

@pytest.mark.asyncio
async def test_configured_resource_limits_enforcement(monkeypatch):
    async with AsyncClient(transport=ASGITransport(app=app), base_url="http://testserver") as ac:
        pfx = uuid.uuid4().hex[:6]
        token, tenant_id = await get_auth_token(ac, f"limits_{pfx}@tenant.com", f"Limits Tenant {pfx}")
        headers = {"Authorization": f"Bearer {token}"}

        # 1. DOCX uncompressed zip size limit
        monkeypatch.setenv("MAX_DOCX_UNCOMPRESSED_SIZE_MB", "0")
        docx_bytes = build_docx_bytes(["A" * 2000])
        res_docx_limit = await ac.post(
            "/documents/upload",
            headers=headers,
            files={"file": ("large_uncompressed.docx", docx_bytes, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")}
        )
        assert res_docx_limit.status_code == 400
        assert "uncompressed size" in res_docx_limit.json()["detail"].lower()

        # 2. Extracted text length limit
        monkeypatch.setenv("MAX_DOCX_UNCOMPRESSED_SIZE_MB", "50")
        monkeypatch.setenv("MAX_EXTRACTED_TEXT_LENGTH", "50")
        res_text_limit = await ac.post(
            "/documents/upload",
            headers=headers,
            files={"file": ("too_long.txt", ("B" * 60).encode("utf-8"), "text/plain")}
        )
        assert res_text_limit.status_code == 400
        assert "text length" in res_text_limit.json()["detail"].lower()

        # 3. Maximum document chunks limit
        monkeypatch.setenv("MAX_EXTRACTED_TEXT_LENGTH", "1000000")
        monkeypatch.setenv("MAX_DOCUMENT_CHUNKS", "1")
        multi_chunk_text = ("First long paragraph that definitely creates multiple chunks in the sliding window. " * 20).encode("utf-8")
        res_chunks_limit = await ac.post(
            "/documents/upload",
            headers=headers,
            files={"file": ("too_many_chunks.txt", multi_chunk_text, "text/plain")}
        )
        assert res_chunks_limit.status_code == 400
        assert "chunk" in res_chunks_limit.json()["detail"].lower()

@pytest.mark.asyncio
async def test_partial_vector_write_failure_compensating_rollback(caplog):
    async with AsyncClient(transport=ASGITransport(app=app), base_url="http://testserver") as ac:
        pfx = uuid.uuid4().hex[:6]
        token, tenant_id = await get_auth_token(ac, f"rollback_{pfx}@tenant.com", f"Rollback Tenant {pfx}")
        headers = {"Authorization": f"Bearer {token}"}
        db = get_db()

        sensitive_vector_err = "QDRANT_AUTH_FAIL: secret_api_key_abc123! at https://internal-qdrant.cluster.local:6333"

        # 1. Test POST /documents/upload failure handling
        with caplog.at_level(logging.ERROR, logger="multitenant_rag.documents"):
            with patch(
                "backend.app.api.documents.upsert_chunks",
                side_effect=RuntimeError(sensitive_vector_err)
            ), patch(
                "backend.app.api.documents.delete_document_chunks",
                wraps=backend.app.api.documents.delete_document_chunks
            ) as mock_delete_chunks:
                res = await ac.post(
                    "/documents/upload",
                    headers=headers,
                    files={"file": ("failing_doc.txt", b"Valid text that will fail vector indexing.", "text/plain")}
                )
                assert res.status_code == 500
                detail = res.json()["detail"]

                # Assert stable, generic HTTP 500 message
                assert detail == "Failed to index document due to an internal processing error."
                # Assert sensitive strings and internal database details never leak in HTTP response
                assert sensitive_vector_err not in res.text
                assert "secret_api_key_abc123" not in res.text
                assert "internal-qdrant.cluster.local" not in res.text

                # Verify tenant-scoped deletion called with correct tenant_id
                mock_delete_chunks.assert_called_once()
                called_tenant = mock_delete_chunks.call_args.kwargs.get("tenant_id") or mock_delete_chunks.call_args.args[0]
                called_doc_id = mock_delete_chunks.call_args.kwargs.get("doc_id") or mock_delete_chunks.call_args.args[1]
                assert called_tenant == tenant_id

                # Verify MongoDB document was deleted by compensating rollback
                deleted_doc = await db.documents.find_one({"_id": ObjectId(called_doc_id), "tenant_id": tenant_id})
                assert deleted_doc is None, "Compensating rollback should have deleted the MongoDB document!"

                # Verify server logs retain diagnostic error info with tenant and doc identifiers
                err_logs = [r for r in caplog.records if r.levelno == logging.ERROR]
                assert any(sensitive_vector_err in r.message for r in err_logs)
                assert any(tenant_id in r.message for r in err_logs)
                assert any(called_doc_id in r.message for r in err_logs)

        # 2. Test legacy POST /documents (raw-text) creation route failure handling
        caplog.clear()
        with caplog.at_level(logging.ERROR, logger="multitenant_rag.documents"):
            with patch(
                "backend.app.api.documents.upsert_chunks",
                side_effect=RuntimeError(sensitive_vector_err)
            ), patch(
                "backend.app.api.documents.delete_document_chunks",
                wraps=backend.app.api.documents.delete_document_chunks
            ) as mock_delete_raw:
                res_raw = await ac.post(
                    "/documents",
                    headers=headers,
                    json={"title": "Raw Fail Doc", "raw_text": "Sample text that fails vector upsert."}
                )
                assert res_raw.status_code == 500
                assert res_raw.json()["detail"] == "Failed to index document due to an internal processing error."
                assert sensitive_vector_err not in res_raw.text
                assert "secret_api_key_abc123" not in res_raw.text

                mock_delete_raw.assert_called_once()
                called_raw_doc_id = mock_delete_raw.call_args.kwargs.get("doc_id") or mock_delete_raw.call_args.args[1]
                deleted_raw = await db.documents.find_one({"_id": ObjectId(called_raw_doc_id), "tenant_id": tenant_id})
                assert deleted_raw is None

@pytest.mark.asyncio
async def test_cleanup_failure_during_compensating_rollback_observable(caplog):
    async with AsyncClient(transport=ASGITransport(app=app), base_url="http://testserver") as ac:
        pfx = uuid.uuid4().hex[:6]
        token, tenant_id = await get_auth_token(ac, f"cleanup_fail_{pfx}@tenant.com", f"Cleanup Fail Tenant {pfx}")
        headers = {"Authorization": f"Bearer {token}"}
        db = get_db()

        sensitive_index_err = "INDEX_ERR: key=SecretIndexKey888 host=internal-vdb:6333"
        sensitive_qdrant_cleanup_err = "QDRANT_CLEANUP_ERR: token=SecretQdrantToken555 host=internal-qdrant:6333"
        sensitive_mongo_cleanup_err = "MONGO_CLEANUP_ERR: mongodb://app_admin:SecretMongoPassword777@mongo-replica:27017"

        # Case 1: Qdrant vector cleanup fails during rollback (MongoDB cleanup succeeds)
        with caplog.at_level(logging.CRITICAL, logger="multitenant_rag.documents"):
            with patch(
                "backend.app.api.documents.upsert_chunks",
                side_effect=RuntimeError(sensitive_index_err)
            ), patch(
                "backend.app.api.documents.delete_document_chunks",
                side_effect=RuntimeError(sensitive_qdrant_cleanup_err)
            ):
                res = await ac.post(
                    "/documents/upload",
                    headers=headers,
                    files={"file": ("cleanup_error.txt", b"Content attempting cleanup", "text/plain")}
                )
                assert res.status_code == 500
                detail = res.json()["detail"]
                assert detail == "Failed to index document due to an internal processing error."
                # Verify sensitive strings never leak into response
                assert sensitive_index_err not in res.text
                assert sensitive_qdrant_cleanup_err not in res.text
                assert "SecretIndexKey888" not in res.text
                assert "SecretQdrantToken555" not in res.text

                # Verify critical logging preserves diagnostic info with tenant_id
                crit_logs = [r for r in caplog.records if r.levelno == logging.CRITICAL]
                assert any(sensitive_qdrant_cleanup_err in r.message for r in crit_logs)
                assert any(tenant_id in r.message for r in crit_logs)

        # Case 2: MongoDB deletion fails during rollback AND Qdrant cleanup is STILL attempted
        # (Verifies both cleanup attempts occur independently and neither blocks the other)
        caplog.clear()
        with caplog.at_level(logging.CRITICAL, logger="multitenant_rag.documents"):
            with patch(
                "backend.app.api.documents.upsert_chunks",
                side_effect=RuntimeError(sensitive_index_err)
            ), patch(
                "backend.app.api.documents.delete_document_chunks",
                wraps=backend.app.api.documents.delete_document_chunks
            ) as mock_qdrant_cleanup, patch(
                "backend.app.api.documents.get_db"
            ) as mock_get_db:
                mock_db = MagicMock()
                mock_db.documents = MagicMock()
                mock_db.documents.insert_one = db.documents.insert_one
                mock_db.documents.delete_one = AsyncMock(side_effect=RuntimeError(sensitive_mongo_cleanup_err))
                mock_get_db.return_value = mock_db

                try:
                    res_mongo_fail = await ac.post(
                        "/documents/upload",
                        headers=headers,
                        files={"file": ("mongo_fail.txt", b"Content with mongo cleanup failure", "text/plain")}
                    )
                    assert res_mongo_fail.status_code == 500
                    detail_mongo = res_mongo_fail.json()["detail"]
                    assert detail_mongo == "Failed to index document due to an internal processing error."
                    # Verify no sensitive leakage
                    assert sensitive_index_err not in res_mongo_fail.text
                    assert sensitive_mongo_cleanup_err not in res_mongo_fail.text
                    assert "SecretMongoPassword777" not in res_mongo_fail.text

                    # Assert BOTH cleanup attempts occurred independently
                    mock_db.documents.delete_one.assert_called_once()
                    mock_qdrant_cleanup.assert_called_once()

                    # Verify critical logging recorded MongoDB cleanup failure
                    crit_logs_mongo = [r for r in caplog.records if r.levelno == logging.CRITICAL]
                    assert any(sensitive_mongo_cleanup_err in r.message for r in crit_logs_mongo)
                    assert any(tenant_id in r.message for r in crit_logs_mongo)
                finally:
                    await db.documents.delete_many({"tenant_id": tenant_id})

        # Case 3: Both MongoDB cleanup AND Qdrant cleanup fail during rollback
        caplog.clear()
        with caplog.at_level(logging.CRITICAL, logger="multitenant_rag.documents"):
            with patch(
                "backend.app.api.documents.upsert_chunks",
                side_effect=RuntimeError(sensitive_index_err)
            ), patch(
                "backend.app.api.documents.delete_document_chunks",
                side_effect=RuntimeError(sensitive_qdrant_cleanup_err)
            ) as mock_qdrant_fail, patch(
                "backend.app.api.documents.get_db"
            ) as mock_get_db:
                mock_db = MagicMock()
                mock_db.documents = MagicMock()
                mock_db.documents.insert_one = db.documents.insert_one
                mock_db.documents.delete_one = AsyncMock(side_effect=RuntimeError(sensitive_mongo_cleanup_err))
                mock_get_db.return_value = mock_db

                try:
                    res_both_fail = await ac.post(
                        "/documents/upload",
                        headers=headers,
                        files={"file": ("both_fail.txt", b"Content where both cleanups fail", "text/plain")}
                    )
                    assert res_both_fail.status_code == 500
                    assert res_both_fail.json()["detail"] == "Failed to index document due to an internal processing error."
                    # None of the sensitive strings appear in the HTTP response
                    assert sensitive_index_err not in res_both_fail.text
                    assert sensitive_mongo_cleanup_err not in res_both_fail.text
                    assert sensitive_qdrant_cleanup_err not in res_both_fail.text

                    # Both cleanup attempts occurred
                    mock_db.documents.delete_one.assert_called_once()
                    mock_qdrant_fail.assert_called_once()

                    # Both cleanup failures logged separately as CRITICAL
                    crit_logs_both = [r for r in caplog.records if r.levelno == logging.CRITICAL]
                    assert any(sensitive_mongo_cleanup_err in r.message for r in crit_logs_both)
                    assert any(sensitive_qdrant_cleanup_err in r.message for r in crit_logs_both)
                    assert any(tenant_id in r.message for r in crit_logs_both)
                finally:
                    await db.documents.delete_many({"tenant_id": tenant_id})
