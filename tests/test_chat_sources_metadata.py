import io
import uuid
from unittest.mock import patch, MagicMock, AsyncMock
import pytest
import docx
from bson import ObjectId
from httpx import AsyncClient, ASGITransport
from qdrant_client.models import PointStruct

from backend.main import app
from backend.app.db.mongo import get_db
from backend.app.db.vector import get_qdrant_client, get_tenant_collection_name, ensure_tenant_collection
from backend.app.services.embedding import get_embedding

def build_pdf_bytes(pages_text: list[str]) -> bytes:
    """Builds a valid multi-page PDF with extractable text."""
    lines = [b'%PDF-1.4\n']
    offsets = {}

    def add_obj(obj_num, data):
        offsets[obj_num] = sum(len(l) for l in lines)
        lines.append(f'{obj_num} 0 obj\n'.encode() + data + b'\nendobj\n')

    total_pages = len(pages_text)
    page_obj_ids = [3 + i * 3 + 2 for i in range(total_pages)]

    add_obj(1, b'<< /Type /Catalog /Pages 2 0 R >>')
    kids_str = ' '.join(f'{pid} 0 R' for pid in page_obj_ids)
    add_obj(2, f'<< /Type /Pages /Kids [{kids_str}] /Count {total_pages} >>'.encode())

    curr = 3
    for text in pages_text:
        cid = curr
        fid = curr + 1
        pid = curr + 2
        curr += 3
        stream_data = f'BT /F1 12 Tf 50 700 Td ({text}) Tj ET'.encode()
        add_obj(cid, f'<< /Length {len(stream_data)} >> stream\n'.encode() + stream_data + b'\nendstream')
        add_obj(fid, b'<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>')
        add_obj(pid, f'<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents {cid} 0 R /Resources << /Font << /F1 {fid} 0 R >> >> >>'.encode())

    startxref = sum(len(l) for l in lines)
    total_objs = curr
    lines.append(f'xref\n0 {total_objs}\n0000000000 65535 f \n'.encode())
    for i in range(1, total_objs):
        lines.append(f'{offsets[i]:010d} 00000 n \n'.encode())
    lines.append(f'trailer << /Root 1 0 R /Size {total_objs} >>\nstartxref\n{startxref}\n%%EOF\n'.encode())
    return b''.join(lines)

def build_docx_bytes(items: list) -> bytes:
    """Builds a valid DOCX with paragraphs and tables."""
    doc = docx.Document()
    for item in items:
        if isinstance(item, str):
            doc.add_paragraph(item)
        elif isinstance(item, list):
            tbl = doc.add_table(rows=len(item), cols=len(item[0]))
            for r_idx, row in enumerate(item):
                for c_idx, cell in enumerate(row):
                    tbl.rows[r_idx].cells[c_idx].text = str(cell)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()

async def get_auth_token(ac: AsyncClient, email: str, tenant_name: str) -> tuple[str, str, str]:
    res = await ac.post("/auth/login", json={"email": email, "password": "password123"})
    if res.status_code == 401:
        res = await ac.post("/auth/register", json={
            "email": email,
            "password": "password123",
            "tenant_name": tenant_name
        })
    assert res.status_code == 200, f"Auth failed for {email}: {res.text}"
    data = res.json()
    return data["access_token"], data["tenant_id"], data["user_id"]

@pytest.fixture(autouse=True)
def mock_llm_calls(monkeypatch):
    """
    Mocks LLM generation calls to ensure test execution does not depend on
    external network calls, while still exercising the full RAG retrieval pipeline.
    """
    monkeypatch.setenv("RAG_SIMILARITY_THRESHOLD", "0.40")
    mock_client = MagicMock()

    def fake_generate_content(*args, **kwargs):
        contents = kwargs.get("contents") or (args[0] if args else "")
        contents_str = str(contents)
        if "User Query:" in contents_str:
            # Step 1: Agent route check; return empty function calls to execute search_kb_func
            return MagicMock(text="Looking up knowledge base.", function_calls=None)
        if "Strict Grounding Rules:" in contents_str:
            # Step 2: Synthesis
            return MagicMock(text="Based on the retrieved documents, here is the answer.", function_calls=None)
        return MagicMock(text="Grounded answer.", function_calls=None)

    mock_client.models.generate_content.side_effect = fake_generate_content
    monkeypatch.setattr("backend.app.services.llm.get_genai_client", lambda: mock_client)
    monkeypatch.setattr("backend.app.services.llm.rewrite_query", AsyncMock(side_effect=lambda history, msg: msg))

@pytest.mark.asyncio
async def test_pdf_sources_metadata_in_chat_and_history():
    """
    Verify PDF page numbers (1-based integer) and filename metadata are
    correctly carried through chat source responses and preserved in saved history.
    """
    async with AsyncClient(transport=ASGITransport(app=app), base_url="http://testserver") as ac:
        pfx = uuid.uuid4().hex[:6]
        token, tenant_id, _ = await get_auth_token(ac, f"pdf_chat_{pfx}@tenant.com", f"PDF Chat Tenant {pfx}")
        headers = {"Authorization": f"Bearer {token}"}

        # Upload a 2-page PDF with distinct text on Page 1 and Page 2
        pdf_bytes = build_pdf_bytes([
            "Project Mercury satellite communications architecture and high-frequency telemetry.",
            "Project Apollo lunar trajectory calculation equations and orbital mechanics."
        ])
        res_upload = await ac.post(
            "/documents/upload",
            headers=headers,
            files={"file": ("space_missions.pdf", pdf_bytes, "application/pdf")},
            data={"title": "Space Missions"}
        )
        assert res_upload.status_code == 201

        # 1. Query for Page 1 topic
        res_chat_p1 = await ac.post("/chat", headers=headers, json={
            "message": "Tell me about Project Mercury satellite communications telemetry"
        })
        assert res_chat_p1.status_code == 200
        data_p1 = res_chat_p1.json()
        conv_id = data_p1["conversation_id"]

        assert len(data_p1["sources"]) > 0, "Expected at least 1 source hit for Page 1"
        source_p1 = [s for s in data_p1["sources"] if "Project Mercury" in s["chunk_text"]][0]
        assert source_p1["document_title"] == "Space Missions"
        assert source_p1["filename"] == "space_missions.pdf"
        assert source_p1["page_number"] == 1
        assert isinstance(source_p1["page_number"], int)

        # 2. Query for Page 2 topic in the same conversation
        res_chat_p2 = await ac.post("/chat", headers=headers, json={
            "conversation_id": conv_id,
            "message": "Explain the Project Apollo lunar trajectory calculation equations and orbital mechanics"
        })
        assert res_chat_p2.status_code == 200
        data_p2 = res_chat_p2.json()

        assert len(data_p2["sources"]) > 0, "Expected at least 1 source hit for Page 2"
        source_p2 = [s for s in data_p2["sources"] if "Project Apollo" in s["chunk_text"]][0]
        assert source_p2["document_title"] == "Space Missions"
        assert source_p2["filename"] == "space_missions.pdf"
        assert source_p2["page_number"] == 2
        assert isinstance(source_p2["page_number"], int)

        # 3. Retrieve conversation history and verify sources are preserved
        res_hist = await ac.get(f"/conversations/{conv_id}/messages", headers=headers)
        assert res_hist.status_code == 200
        history = res_hist.json()
        assert history["conversation_id"] == conv_id

        asst_msgs = [m for m in history["messages"] if m["role"] == "assistant"]
        assert len(asst_msgs) == 2

        # Check first assistant response history
        h_src_p1 = [s for s in asst_msgs[0]["sources"] if "Project Mercury" in s["chunk_text"]][0]
        assert h_src_p1["filename"] == "space_missions.pdf"
        assert h_src_p1["page_number"] == 1

        # Check second assistant response history
        h_src_p2 = [s for s in asst_msgs[1]["sources"] if "Project Apollo" in s["chunk_text"]][0]
        assert h_src_p2["filename"] == "space_missions.pdf"
        assert h_src_p2["page_number"] == 2

@pytest.mark.asyncio
async def test_docx_txt_and_raw_text_null_page_metadata():
    """
    Verify DOCX, TXT, and raw-text uploads carry filename metadata while strictly
    setting page_number=None, both in chat responses and saved history.
    """
    async with AsyncClient(transport=ASGITransport(app=app), base_url="http://testserver") as ac:
        pfx = uuid.uuid4().hex[:6]
        token, tenant_id, _ = await get_auth_token(ac, f"null_page_{pfx}@tenant.com", f"Null Page Tenant {pfx}")
        headers = {"Authorization": f"Bearer {token}"}

        # 1. Upload DOCX
        docx_bytes = build_docx_bytes([
            "Section 1: Enterprise Network Architecture Guidelines",
            [["Subnet", "VLAN", "CIDR"], ["Prod-A", "100", "10.0.1.0/24"], ["Prod-B", "200", "10.0.2.0/24"]]
        ])
        res_docx = await ac.post(
            "/documents/upload",
            headers=headers,
            files={"file": ("network_guidelines.docx", docx_bytes, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")}
        )
        assert res_docx.status_code == 201

        # 2. Upload TXT
        txt_bytes = "Database backup policy: automated daily snapshots are preserved for 30 days.".encode("utf-8")
        res_txt = await ac.post(
            "/documents/upload",
            headers=headers,
            files={"file": ("backup_policy.txt", txt_bytes, "text/plain")}
        )
        assert res_txt.status_code == 201

        # 3. Upload Raw-text legacy payload
        res_raw = await ac.post(
            "/documents",
            headers=headers,
            json={
                "title": "Raw System Architecture",
                "raw_text": "Monolithic core services are decomposed into isolated multi-tenant microservices."
            }
        )
        assert res_raw.status_code == 201

        # Query DOCX
        res_chat_docx = await ac.post("/chat", headers=headers, json={
            "message": "What is the VLAN for Prod-A in network guidelines?"
        })
        assert res_chat_docx.status_code == 200
        sources_docx = res_chat_docx.json()["sources"]
        assert len(sources_docx) > 0
        assert sources_docx[0]["filename"] == "network_guidelines.docx"
        assert sources_docx[0]["page_number"] is None

        # Query TXT
        res_chat_txt = await ac.post("/chat", headers=headers, json={
            "message": "How long are database backup snapshots preserved?"
        })
        assert res_chat_txt.status_code == 200
        sources_txt = res_chat_txt.json()["sources"]
        assert len(sources_txt) > 0
        assert sources_txt[0]["filename"] == "backup_policy.txt"
        assert sources_txt[0]["page_number"] is None

        # Query Raw-text
        res_chat_raw = await ac.post("/chat", headers=headers, json={
            "message": "How are monolithic core services decomposed?"
        })
        assert res_chat_raw.status_code == 200
        sources_raw = res_chat_raw.json()["sources"]
        assert len(sources_raw) > 0
        assert sources_raw[0]["filename"] is None
        assert sources_raw[0]["page_number"] is None

        # Verify saved conversation history for the raw-text chat
        conv_id = res_chat_raw.json()["conversation_id"]
        res_hist = await ac.get(f"/conversations/{conv_id}/messages", headers=headers)
        assert res_hist.status_code == 200
        hist_msg = [m for m in res_hist.json()["messages"] if m["role"] == "assistant"][0]
        assert hist_msg["sources"][0]["filename"] is None
        assert hist_msg["sources"][0]["page_number"] is None

@pytest.mark.asyncio
async def test_older_payload_compatibility_and_null_metadata():
    """
    Verify backward compatibility with older Qdrant payloads and MongoDB messages
    that lack 'filename' and 'page_number' metadata.
    """
    async with AsyncClient(transport=ASGITransport(app=app), base_url="http://testserver") as ac:
        pfx = uuid.uuid4().hex[:6]
        token, tenant_id, user_id = await get_auth_token(ac, f"legacy_{pfx}@tenant.com", f"Legacy Tenant {pfx}")
        headers = {"Authorization": f"Bearer {token}"}
        db = get_db()

        # 1. Seed Qdrant directly with older payload format (no page_number, no filename)
        client = get_qdrant_client()
        coll_name = get_tenant_collection_name(tenant_id)
        vec_sample = get_embedding("Legacy archival specification document chunk.")
        await ensure_tenant_collection(tenant_id, vector_size=len(vec_sample))

        legacy_doc_id = ObjectId()
        await db.documents.insert_one({
            "_id": legacy_doc_id,
            "tenant_id": tenant_id,
            "title": "Legacy Archival Spec",
            "raw_text": "Legacy archival specification document chunk.",
            "chunk_count": 1,
            "uploaded_at": "2025-01-01T00:00:00Z"
        })

        old_point_id = str(uuid.uuid4())
        await client.upsert(
            collection_name=coll_name,
            points=[
                PointStruct(
                    id=old_point_id,
                    vector=vec_sample,
                    payload={
                        "document_id": str(legacy_doc_id),
                        "tenant_id": str(tenant_id),
                        "document_title": "Legacy Archival Spec",
                        "chunk_index": 0,
                        "text": "Legacy archival specification document chunk."
                        # Note: 'filename' and 'page_number' intentionally omitted
                    }
                )
            ]
        )

        res_chat = await ac.post("/chat", headers=headers, json={
            "message": "Tell me about the legacy archival specification"
        })
        assert res_chat.status_code == 200
        data_chat = res_chat.json()
        assert len(data_chat["sources"]) > 0
        src = data_chat["sources"][0]
        assert src["document_title"] == "Legacy Archival Spec"
        assert src["filename"] is None
        assert src["page_number"] is None

        # 2. Seed MongoDB directly with an older message format (sources without page_number/filename)
        conv_obj_id = ObjectId()
        await db.conversations.insert_one({
            "_id": conv_obj_id,
            "tenant_id": tenant_id,
            "user_id": user_id,
            "created_at": "2025-01-01T00:00:00Z"
        })
        await db.messages.insert_one({
            "conversation_id": str(conv_obj_id),
            "tenant_id": tenant_id,
            "user_id": user_id,
            "role": "assistant",
            "content": "Historic answer from legacy system.",
            "sources": [
                {
                    "document_title": "Historic Document",
                    "chunk_text": "Historic text without modern metadata.",
                    "score": 0.88
                    # 'filename' and 'page_number' omitted
                }
            ],
            "created_at": "2025-01-01T00:01:00Z"
        })

        res_hist = await ac.get(f"/conversations/{str(conv_obj_id)}/messages", headers=headers)
        assert res_hist.status_code == 200
        hist_data = res_hist.json()
        assert len(hist_data["messages"]) == 1
        hist_src = hist_data["messages"][0]["sources"][0]
        assert hist_src["document_title"] == "Historic Document"
        assert hist_src["filename"] is None
        assert hist_src["page_number"] is None

@pytest.mark.asyncio
async def test_cross_tenant_exclusion_and_zero_sources_refusal():
    """
    Verify cross-tenant documents are never leaked into chat sources, and
    refused/unsupported queries return zero sources.
    """
    async with AsyncClient(transport=ASGITransport(app=app), base_url="http://testserver") as ac:
        pfx = uuid.uuid4().hex[:6]
        token_a, tenant_a, _ = await get_auth_token(ac, f"tenant_a_{pfx}@corp.com", f"Tenant A {pfx}")
        headers_a = {"Authorization": f"Bearer {token_a}"}

        token_b, tenant_b, _ = await get_auth_token(ac, f"tenant_b_{pfx}@corp.com", f"Tenant B {pfx}")
        headers_b = {"Authorization": f"Bearer {token_b}"}

        # Tenant A uploads a proprietary PDF
        pdf_a = build_pdf_bytes([
            "Confidential Formula X: The activation catalyst is compound QK-8902."
        ])
        res_up_a = await ac.post(
            "/documents/upload",
            headers=headers_a,
            files={"file": ("secret_formula.pdf", pdf_a, "application/pdf")},
            data={"title": "Secret Formula"}
        )
        assert res_up_a.status_code == 201

        # 1. Tenant B asks about Tenant A's document content -> Must refuse with ZERO sources
        res_chat_b = await ac.post("/chat", headers=headers_b, json={
            "message": "What is the activation catalyst compound for Confidential Formula X?"
        })
        assert res_chat_b.status_code == 200
        data_b = res_chat_b.json()
        assert "I couldn't find this information in your documents." in data_b["answer"]
        assert data_b["sources"] == [], "Cross-tenant retrieval must return zero sources"

        # Verify Tenant B's saved history contains zero sources
        res_hist_b = await ac.get(f"/conversations/{data_b['conversation_id']}/messages", headers=headers_b)
        assert res_hist_b.status_code == 200
        b_asst_msg = [m for m in res_hist_b.json()["messages"] if m["role"] == "assistant"][0]
        assert b_asst_msg["sources"] == []

        # 2. Tenant A asks an unrelated question -> Must refuse with ZERO sources
        res_unrelated = await ac.post("/chat", headers=headers_a, json={
            "message": "What is the average rainfall in the Amazon rainforest in November?"
        })
        assert res_unrelated.status_code == 200
        data_unrelated = res_unrelated.json()
        assert "I couldn't find this information in your documents." in data_unrelated["answer"]
        assert data_unrelated["sources"] == [], "Unrelated query must return zero sources"
